#!/usr/bin/env python3
"""Compile every Sluggers Stat Tracker .xlsx into league-wide totals.

Reads the tracker's Game Info / Stats / Pitching sheets. Counting stats are
summed across games; rates are recomputed from those totals.

Player identity starts as (team, name, primary position). After all games are
in, batting rows for the same name on a team are compressed into one row when
their combined games do not exceed that team's games (a player who moved).
Two of the same name who overlap stay split by position.

Pitchers are one row per (team, name), except names that more than one player
on a team carried in the same game (every Mii of one color shares a name). Those
are split by the primary position of the lineup row each one pitched from, for
example "Brown Mii (M) [P]" and "Brown Mii (M) [SS]".
"""

from __future__ import annotations

import argparse
import re
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

SKIP_NAMES = {"stat_template_do_not_remove.xlsx"}
FILENAME_TS = re.compile(r"(\d{4}-\d{2}-\d{2} \d{2}-\d{2}-\d{2})")

BATTING_SUM = [
    "At-Bats",
    "Plate Appearances",
    "Runs",
    "Hits",
    "RBI",
    "Strikeouts",
    "Walks",
    "Hit By Pitch",
    "Singles",
    "Doubles",
    "Triples",
    "Home Runs",
    "1HR",
    "2HR",
    "3HR",
    "Grand Slams",
    "ITP Home Runs",
    "Total Bases",
    "Sac Flys",
    "Double Plays Hit Into",
    "Triple Plays Hit Into",
    "Buddy Jumps Hit Into",
    "Star Swings",
    "Star Hits",
    "Stars Used",
    "Stolen Bases",
    "Caught Stealing",
    "Steal Attempts",
    "Putouts",
    "Assists",
    "Buddy Jump Putouts",
    "Buddy Jump Attempts",
    "Double Plays",
    "Triple Plays",
    "Bobbles",
    "Star Bases Helper",
]

PITCHING_SUM = [
    "Batters Faced",
    "Innings Pitched",
    "Pitches",
    "Strikes",
    "Balls",
    "Strikeouts",
    "Walks",
    "Bean Balls",
    "Hits Allowed",
    "Runs Allowed",
    "Singles Allowed",
    "Doubles Allowed",
    "Triples Allowed",
    "HR Allowed",
    "Earned Runs",
    "Inherited Runs",
    "Star Pitches",
    "Stars Used",
    "Pickoffs",
    "Pickoff Attempts",
    "At-Bats Against Helper",
    "Total Bases Allowed Helper",
]

# Output columns as (header, field). Helpers behind the rates (Star Bases, At-Bats Against,
# Total Bases Allowed) stay out; every play the tracker counts is in.
BATTING_COLUMNS = [
    ("Team", "team"), ("Player", "player"), ("Pos", "position"), ("G", "games"),
    ("PA", "Plate Appearances"), ("AB", "At-Bats"), ("R", "Runs"), ("H", "Hits"),
    ("1B", "Singles"), ("2B", "Doubles"), ("3B", "Triples"), ("HR", "Home Runs"),
    ("Solo HR", "1HR"), ("2-Run HR", "2HR"), ("3-Run HR", "3HR"), ("GS", "Grand Slams"), ("ITP HR", "ITP Home Runs"),
    ("RBI", "RBI"), ("BB", "Walks"), ("HBP", "Hit By Pitch"), ("SO", "Strikeouts"), ("TB", "Total Bases"),
    ("AVG", "Batting Average"), ("OBP", "On Base %"), ("SLG", "Slug %"), ("OPS", "On Base + Slug"), ("Star SLG", "Star Slug %"),
    ("SF", "Sac Flys"), ("GIDP", "Double Plays Hit Into"), ("TP into", "Triple Plays Hit Into"), ("BJ into", "Buddy Jumps Hit Into"),
    ("SB", "Stolen Bases"), ("CS", "Caught Stealing"), ("SB Att", "Steal Attempts"),
    ("Star Swings", "Star Swings"), ("Star Hits", "Star Hits"), ("Stars Used", "Stars Used"),
    ("PO", "Putouts"), ("A", "Assists"), ("DP", "Double Plays"), ("TP", "Triple Plays"),
    ("BJ PO", "Buddy Jump Putouts"), ("BJ Att", "Buddy Jump Attempts"), ("Bobbles", "Bobbles"),
]
PITCHING_COLUMNS = [
    ("Team", "team"), ("Player", "player"), ("G", "games"), ("IP", "Innings Pitched"), ("BF", "Batters Faced"),
    ("Pitches", "Pitches"), ("Strikes", "Strikes"), ("Balls", "Balls"),
    ("K", "Strikeouts"), ("BB", "Walks"), ("HBP", "Bean Balls"),
    ("H", "Hits Allowed"), ("1B Alw", "Singles Allowed"), ("2B Alw", "Doubles Allowed"), ("3B Alw", "Triples Allowed"), ("HR", "HR Allowed"),
    ("R", "Runs Allowed"), ("ER", "Earned Runs"), ("Inh R", "Inherited Runs"),
    ("ERA-7", "ERA-7"), ("ERA-9", "ERA-9"), ("WHIP", "WHIP"),
    ("BAA", "BA Against"), ("OBP-A", "OB% Against"), ("SLG-A", "SLG Against"), ("OPS-A", "OPS Against"),
    ("Star Pitches", "Star Pitches"), ("Stars Used", "Stars Used"), ("Pickoffs", "Pickoffs"), ("PK Att", "Pickoff Attempts"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
ALT_FILL = PatternFill("solid", fgColor="D6EAF8")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)


def _num(value) -> float:
    if value is None or value == "" or value == "INF":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _rate(num: float, den: float) -> float | None:
    if den <= 0:
        return None
    return num / den


def _headers(ws) -> dict[str, int]:
    out = {}
    for col in range(1, (ws.max_column or 1) + 1):
        header = ws.cell(1, col).value
        if header:
            out[str(header).strip()] = col
    return out


def _row_map(ws, row: int, headers: dict[str, int]) -> dict[str, object]:
    return {name: ws.cell(row, col).value for name, col in headers.items()}


def _primary_position(raw) -> str:
    text = str(raw or "").strip()
    if not text or text.upper() == "N/A":
        return ""
    return text.split(",")[0].strip()


def _is_total_row(player: str, position: str, team: str) -> bool:
    name = str(player or "").strip()
    pos = str(position or "").strip()
    team_name = str(team or "").strip()
    if not name:
        return True
    if name.lower().endswith(" totals"):
        return True
    if pos.upper() == "N/A" and name == team_name:
        return True
    if name == team_name and pos.upper() in {"", "N/A"}:
        return True
    return False


def _parse_game_info(ws) -> dict[str, object]:
    away = ws["L11"].value
    home = ws["P11"].value
    away_score = _num(ws["I12"].value)
    home_score = _num(ws["P12"].value)
    stadium = ws["I14"].value
    innings = ws["M15"].value
    started_late = bool(ws["L9"].value)
    quit_early = bool(ws["L10"].value)
    return {
        "away": str(away or "").strip(),
        "home": str(home or "").strip(),
        "away_score": int(away_score),
        "home_score": int(home_score),
        "stadium": str(stadium or "").strip(),
        "innings": str(innings or "").strip(),
        "started_late": started_late,
        "quit_early": quit_early,
    }


def _timestamp_from_name(path: Path) -> str:
    match = FILENAME_TS.search(path.stem)
    return match.group(1) if match else ""


def iter_workbooks(folder: Path) -> list[Path]:
    files = []
    for path in sorted(folder.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        if path.name.lower() in SKIP_NAMES:
            continue
        if path.name.lower().startswith("league_"):
            continue
        files.append(path)
    return files


def read_game(path: Path) -> dict[str, object]:
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        if "Stats" not in wb.sheetnames:
            raise ValueError(f"{path.name}: no Stats sheet")
        info = _parse_game_info(wb["Game Info"]) if "Game Info" in wb.sheetnames else {}
        stats_ws = wb["Stats"]
        stats_headers = _headers(stats_ws)
        if "Player" not in stats_headers:
            raise ValueError(f"{path.name}: Stats sheet missing Player column")

        batters = []
        current_team = ""
        for row in range(2, (stats_ws.max_row or 1) + 1):
            mapped = _row_map(stats_ws, row, stats_headers)
            team_cell = str(mapped.get("Team") or "").strip()
            if team_cell:
                current_team = team_cell
            player = str(mapped.get("Player") or "").strip()
            position = str(mapped.get("Position") or "").strip()
            if _is_total_row(player, position, current_team):
                continue
            if not player or not current_team:
                continue
            counts = {key: _num(mapped.get(key)) for key in BATTING_SUM if key in stats_headers}
            for key in BATTING_SUM:
                counts.setdefault(key, 0.0)
            batters.append(
                {
                    "team": current_team,
                    "player": player,
                    "position": _primary_position(position),
                    "positions": position,
                    "counts": counts,
                }
            )

        pitchers = []
        if "Pitching" in wb.sheetnames:
            pitch_ws = wb["Pitching"]
            pitch_headers = _headers(pitch_ws)
            for row in range(2, (pitch_ws.max_row or 1) + 1):
                mapped = _row_map(pitch_ws, row, pitch_headers)
                player = str(mapped.get("Player") or "").strip()
                team = str(mapped.get("Team") or "").strip()
                if _is_total_row(player, "", team):
                    continue
                if not player or not team:
                    continue
                counts = {key: _num(mapped.get(key)) for key in PITCHING_SUM if key in pitch_headers}
                for key in PITCHING_SUM:
                    counts.setdefault(key, 0.0)
                pitchers.append({"team": team, "player": player, "counts": counts})
    finally:
        wb.close()

    return {"path": path, "info": info, "batters": batters, "pitchers": pitchers}


def _games_by_team(games: list[dict[str, object]]) -> dict[str, int]:
    counts: dict[str, int] = defaultdict(int)
    for game in games:
        for team in {batter["team"] for batter in game["batters"] if batter.get("team")}:
            counts[team] += 1
    return counts


def _compress_batting(
    batting: dict[tuple[str, str, str], dict],
    team_games: dict[str, int],
) -> dict[tuple[str, str, str], dict]:
    grouped: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for slot in batting.values():
        grouped[(slot["team"], slot["player"])].append(slot)
    compressed: dict[tuple[str, str, str], dict] = {}
    for (team, player), slots in grouped.items():
        total = sum(slot["games"] for slot in slots)
        cap = team_games.get(team, 0)
        if len(slots) == 1 or (cap and total > cap):
            for slot in slots:
                compressed[(slot["team"], slot["player"], slot["position"])] = slot
            continue
        counts: dict[str, float] = defaultdict(float)
        for slot in slots:
            for stat, value in slot["counts"].items():
                counts[stat] += value
        positions = [
            pos
            for _, pos in sorted(
                ((slot["games"], slot["position"]) for slot in slots),
                key=lambda item: (-item[0], item[1] or ""),
            )
            if pos
        ]
        position = "/".join(dict.fromkeys(positions))
        compressed[(team, player, position)] = {
            "team": team,
            "player": player,
            "position": position,
            "games": total,
            "counts": counts,
        }
    return compressed


def _shared_names(games: list[dict[str, object]]) -> set[tuple[str, str]]:
    """(team, name) pairs that more than one player on a team carried in the same game."""
    shared: set[tuple[str, str]] = set()
    for game in games:
        for rows in (game["batters"], game["pitchers"]):
            counts = Counter((row["team"], row["player"]) for row in rows)
            shared.update(key for key, n in counts.items() if n > 1)
    return shared


def _assign_pitcher_positions(batters: list[dict], pitchers: list[dict]) -> None:
    """Give each pitcher in one game the primary position of the lineup row they pitched from.

    The Pitching sheet has no position column, and two Miis of one color can both pitch in a
    game under the same name. The Stats sheet lists each player's positions in order: the one
    who started on the mound reads "P, 1B" and a reliever reads "SS, P". Pitching rows are in
    the order the pitchers came in, so the starter comes first, then relievers in sheet order.
    """
    pitched_from: dict[tuple[str, str], list[list[str]]] = defaultdict(list)
    for batter in batters:
        positions = [p.strip() for p in str(batter.get("positions") or "").split(",") if p.strip()]
        if "P" in positions:
            pitched_from[(batter["team"], batter["player"])].append(positions)
    for rows in pitched_from.values():
        rows.sort(key=lambda positions: positions[0] != "P")
    seen: dict[tuple[str, str], int] = defaultdict(int)
    for pitcher in pitchers:
        key = (pitcher["team"], pitcher["player"])
        rows = pitched_from.get(key, [])
        index = seen[key]
        seen[key] += 1
        pitcher["position"] = rows[index][0] if index < len(rows) else ""


def compile_league(games: list[dict[str, object]]) -> dict[str, object]:
    batting: dict[tuple[str, str, str], dict] = {}
    pitching: dict[tuple[str, str, str], dict] = {}
    shared = _shared_names(games)
    standings: dict[str, dict] = defaultdict(lambda: {
        "team": "", "games": 0, "wins": 0, "losses": 0, "ties": 0,
        "runs_for": 0, "runs_against": 0,
    })
    game_rows = []

    for game in games:
        info = game["info"]
        path: Path = game["path"]
        away = info.get("away") or ""
        home = info.get("home") or ""
        away_score = int(info.get("away_score") or 0)
        home_score = int(info.get("home_score") or 0)
        incomplete = bool(info.get("quit_early") or info.get("started_late"))

        if away and home:
            for team, scored, allowed, won, lost, tied in (
                (away, away_score, home_score, away_score > home_score, away_score < home_score, away_score == home_score),
                (home, home_score, away_score, home_score > away_score, home_score < away_score, home_score == away_score),
            ):
                row = standings[team]
                row["team"] = team
                row["games"] += 1
                row["runs_for"] += scored
                row["runs_against"] += allowed
                if tied:
                    row["ties"] += 1
                elif won:
                    row["wins"] += 1
                else:
                    row["losses"] += 1

        game_rows.append({
            "file": path.name,
            "played": _timestamp_from_name(path),
            "away": away,
            "home": home,
            "away_score": away_score,
            "home_score": home_score,
            "stadium": info.get("stadium") or "",
            "innings": info.get("innings") or "",
            "incomplete": incomplete,
        })

        for batter in game["batters"]:
            key = (batter["team"], batter["player"], batter["position"])
            slot = batting.setdefault(key, {
                "team": batter["team"],
                "player": batter["player"],
                "position": batter["position"],
                "games": 0,
                "counts": defaultdict(float),
            })
            slot["games"] += 1
            for stat, value in batter["counts"].items():
                slot["counts"][stat] += value

        _assign_pitcher_positions(game["batters"], game["pitchers"])
        for pitcher in game["pitchers"]:
            name = (pitcher["team"], pitcher["player"])
            position = pitcher.get("position", "") if name in shared else ""
            key = (pitcher["team"], pitcher["player"], position)
            slot = pitching.setdefault(key, {
                "team": pitcher["team"],
                "player": pitcher["player"],
                "position": position,
                "games": 0,
                "counts": defaultdict(float),
            })
            slot["games"] += 1
            for stat, value in pitcher["counts"].items():
                slot["counts"][stat] += value

    # A shared name that split into several pitchers gets its position in the label.
    rows_per_name = Counter((slot["team"], slot["player"]) for slot in pitching.values())
    for slot in pitching.values():
        if rows_per_name[(slot["team"], slot["player"])] > 1:
            slot["player"] = f'{slot["player"]} [{slot["position"] or "?"}]'

    batting = _compress_batting(batting, _games_by_team(games))

    batting_rows = []
    for slot in batting.values():
        c = slot["counts"]
        ab, pa, h = c["At-Bats"], c["Plate Appearances"], c["Hits"]
        bb, hbp, tb = c["Walks"], c["Hit By Pitch"], c["Total Bases"]
        batting_rows.append({
            **{k: slot[k] for k in ("team", "player", "position", "games")},
            **{k: c[k] for k in BATTING_SUM},
            "Batting Average": _rate(h, ab),
            "On Base %": _rate(h + bb + hbp, pa),
            "Slug %": _rate(tb, ab),
            "On Base + Slug": None,
            "Star Slug %": _rate(c["Star Bases Helper"], ab),
        })
        obp = batting_rows[-1]["On Base %"]
        slg = batting_rows[-1]["Slug %"]
        batting_rows[-1]["On Base + Slug"] = None if obp is None or slg is None else obp + slg

    pitching_rows = []
    for slot in pitching.values():
        c = slot["counts"]
        ip = c["Innings Pitched"]
        er, ha, bb, hbp = c["Earned Runs"], c["Hits Allowed"], c["Walks"], c["Bean Balls"]
        bf, aba, tba = c["Batters Faced"], c["At-Bats Against Helper"], c["Total Bases Allowed Helper"]
        pitching_rows.append({
            **{k: slot[k] for k in ("team", "player", "games")},
            **{k: c[k] for k in PITCHING_SUM},
            "ERA-7": _rate(er * 7.0, ip),
            "ERA-9": _rate(er * 9.0, ip),
            "WHIP": _rate(ha + bb, ip),
            "BA Against": _rate(ha, aba),
            "OB% Against": _rate(ha + bb + hbp, bf),
            "SLG Against": _rate(tba, aba),
            "OPS Against": None,
        })
        oba = pitching_rows[-1]["OB% Against"]
        slga = pitching_rows[-1]["SLG Against"]
        pitching_rows[-1]["OPS Against"] = None if oba is None or slga is None else oba + slga

    standings_rows = []
    for row in standings.values():
        g = row["games"]
        win_pct = _rate(row["wins"] + 0.5 * row["ties"], g)
        standings_rows.append({**row, "win_pct": win_pct, "run_diff": row["runs_for"] - row["runs_against"]})

    standings_rows.sort(key=lambda r: (-(r["win_pct"] or 0), -r["run_diff"], -r["runs_for"], r["team"]))
    batting_rows.sort(key=lambda r: (-(r["On Base %"] or 0), -(r["On Base + Slug"] or 0), -r["Plate Appearances"], r["player"]))
    pitching_rows.sort(key=lambda r: (-r["Innings Pitched"], (r["ERA-7"] is None, r["ERA-7"] if r["ERA-7"] is not None else 0), r["player"]))
    game_rows.sort(key=lambda r: r["played"] or r["file"])

    return {
        "games": game_rows,
        "standings": standings_rows,
        "batting": batting_rows,
        "pitching": pitching_rows,
    }


def _style_header(ws, width: int, freeze: str = "A2") -> None:
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(width)}{ws.max_row}"
    for col in range(1, width + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _format_sets(headers: list[str], pct: set[str], rate: set[str], text: set[str]) -> tuple[set[int], set[int], set[int]]:
    """Column numbers (1-based) for 0.000 rates, whole numbers and 0.00 rates, chosen by header name."""
    pct_cols = {i for i, h in enumerate(headers, 1) if h in pct}
    rate_cols = {i for i, h in enumerate(headers, 1) if h in rate}
    int_cols = {i for i, h in enumerate(headers, 1) if h not in pct | rate | text}
    return pct_cols, int_cols, rate_cols


def _write_sheet(ws, headers: list[str], rows: list[list], pct_cols: set[int], int_cols: set[int], rate_cols: set[int], freeze: str = "A2") -> None:
    for col, header in enumerate(headers, 1):
        ws.cell(1, col, header)
    for r_i, row in enumerate(rows, 2):
        for c_i, value in enumerate(row, 1):
            cell = ws.cell(r_i, c_i, value)
            cell.border = THIN
            if r_i % 2 == 0:
                cell.fill = ALT_FILL
            if value is None:
                cell.value = "—"
            elif c_i in pct_cols and isinstance(value, float):
                cell.number_format = "0.000"
            elif c_i in rate_cols and isinstance(value, float):
                cell.number_format = "0.00"
            elif c_i in int_cols and isinstance(value, (int, float)):
                cell.number_format = "0"
                cell.value = int(round(value))
            elif isinstance(value, float):
                cell.number_format = "0.00"
    _style_header(ws, len(headers), freeze)
    for col in range(1, len(headers) + 1):
        header = headers[col - 1]
        # Text columns grow to their longest value so names like "Rosalina (Light Blue, F)" are not cut off.
        longest = max((len(row[col - 1]) for row in rows if isinstance(row[col - 1], str)), default=0)
        width = max(len(header) + 2, longest + 2, 12)
        ws.column_dimensions[get_column_letter(col)].width = min(width, 40)


def write_workbook(compiled: dict[str, object], dest: Path, source_count: int) -> None:
    wb = Workbook()

    cover = wb.active
    cover.title = "Summary"
    cover["A1"] = "Mario Super Sluggers — League compile"
    cover["A1"].font = Font(bold=True, size=16)
    cover["A3"] = "Games"
    cover["B3"] = source_count
    cover["A4"] = "Teams"
    cover["B4"] = len(compiled["standings"])
    cover["A5"] = "Batters (team + name + position)"
    cover["B5"] = len(compiled["batting"])
    cover["A6"] = "Pitchers (team + name; shared Mii names by position)"
    cover["B6"] = len(compiled["pitching"])
    cover["A8"] = (
        "Rates are recomputed from summed counting stats. "
        "Position movers with the same name are one batting row; overlapping duplicates stay split. "
        "Miis who share a name and pitch are split by the position they pitched from."
    )
    cover.column_dimensions["A"].width = 42
    cover.column_dimensions["B"].width = 14

    stand_headers = ["Team", "G", "W", "L", "T", "Win %", "RS", "RA", "Diff"]
    stand_rows = [
        [r["team"], r["games"], r["wins"], r["losses"], r["ties"], r["win_pct"], r["runs_for"], r["runs_against"], r["run_diff"]]
        for r in compiled["standings"]
    ]
    stand = wb.create_sheet("Standings")
    _write_sheet(stand, stand_headers, stand_rows, {6}, {2, 3, 4, 5, 7, 8, 9}, set())

    # Every play the tracker counts gets a column; each header sits next to the stat it shows.
    bat_cols = BATTING_COLUMNS
    bat_headers = [h for h, _ in bat_cols]
    bat_rows = [[r[key] for _, key in bat_cols] for r in compiled["batting"]]
    batting_sheet = wb.create_sheet("Batting")
    _write_sheet(
        batting_sheet,
        bat_headers,
        bat_rows,
        *_format_sets(bat_headers, pct={"AVG", "OBP", "SLG", "OPS", "Star SLG"}, rate=set(), text={"Team", "Player", "Pos"}),
        freeze="C2",
    )

    pit_cols = PITCHING_COLUMNS
    pit_headers = [h for h, _ in pit_cols]
    pit_rows = [[r[key] for _, key in pit_cols] for r in compiled["pitching"]]
    pitching_sheet = wb.create_sheet("Pitching")
    _write_sheet(
        pitching_sheet,
        pit_headers,
        pit_rows,
        *_format_sets(pit_headers, pct={"BAA", "OBP-A", "SLG-A", "OPS-A"}, rate={"IP", "ERA-7", "ERA-9", "WHIP"}, text={"Team", "Player"}),
        freeze="C2",
    )

    game_headers = ["Played", "Away", "A", "Home", "H", "Stadium", "Innings", "Incomplete", "File"]
    game_rows = [
        [g["played"], g["away"], g["away_score"], g["home"], g["home_score"], g["stadium"], g["innings"], "yes" if g["incomplete"] else "", g["file"]]
        for g in compiled["games"]
    ]
    games_sheet = wb.create_sheet("Games")
    _write_sheet(games_sheet, game_headers, game_rows, set(), {3, 5}, set())
    games_sheet.column_dimensions["I"].width = 46

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)


def default_input_dir() -> Path | None:
    """Prefer the current directory's game dump, not this machine's tracker path."""
    cwd = Path.cwd()
    for candidate in (cwd / "output", cwd):
        if candidate.is_dir() and iter_workbooks(candidate):
            return candidate
    here = Path(__file__).resolve().parent
    tracker_output = here / "output"
    if (here / "stat_tracker.py").exists() and tracker_output.is_dir() and iter_workbooks(tracker_output):
        return tracker_output
    return None


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="sluggers-league",
        description="Compile Sluggers Stat Tracker .xlsx dumps into league-wide standings, batting, and pitching.",
    )
    parser.add_argument(
        "input",
        nargs="?",
        help="Folder of tracker .xlsx files (default: ./output, then ., then the tracker output folder if present)",
    )
    parser.add_argument("-o", "--output", help="Output workbook path (default: <input>/league_compiled.xlsx)")
    args = parser.parse_args(argv)

    if args.input:
        folder = Path(args.input).expanduser().resolve()
    else:
        folder = default_input_dir()
        if folder is None:
            parser.error("No tracker .xlsx files found. Pass the folder: sluggers-league \"C:\\path\\to\\games\"")
    if not folder.is_dir():
        raise SystemExit(f"Not a folder: {folder}")
    files = iter_workbooks(folder)
    if not files:
        raise SystemExit(f"No tracker .xlsx files in {folder}")

    games = []
    errors = []
    for path in files:
        try:
            games.append(read_game(path))
        except Exception as exc:
            errors.append(f"{path.name}: {exc}")
    if not games:
        raise SystemExit("No readable games.\n" + "\n".join(errors))

    compiled = compile_league(games)
    dest = Path(args.output).expanduser().resolve() if args.output else folder / "league_compiled.xlsx"
    write_workbook(compiled, dest, len(games))

    print(f"Games: {len(games)}")
    print(f"Teams: {len(compiled['standings'])}")
    print(f"Batters: {len(compiled['batting'])}")
    print(f"Pitchers: {len(compiled['pitching'])}")
    print(f"Wrote {dest}")
    if errors:
        print("Skipped:")
        for line in errors:
            print(" ", line)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
